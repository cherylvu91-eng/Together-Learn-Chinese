"""
Script chuyển đổi file Excel từ vựng sang file JSON cho website tra cứu HSK.

CÁCH DÙNG (không cần biết code):
1. Mở Terminal / Command Prompt tại thư mục chứa file này.
2. Chạy lệnh:  python convert_excel_to_json.py
   (Mặc định script sẽ tìm file "vocab_data.xlsx" và tạo ra "data/vocab.json")
3. Nếu muốn dùng tên file khác, chạy:
   python convert_excel_to_json.py "ten_file_cua_ban.xlsx"

YÊU CẦU FILE EXCEL:
- Dòng đầu tiên là tiêu đề cột (không được đổi tên cột, chỉ được thêm dòng dữ liệu).
- Các cột cần có (đúng chính tả, có thể đổi thứ tự các cột thoải mái):
    Từ/Cụm từ | Pinyin | Hán Việt | Nghĩa tiếng Việt | Loại từ | Ví dụ tiếng Trung | Cấp độ HSK | Audio
- Cột "Cấp độ HSK" chỉ nhập số: 1, 2, 3, 4, 5, 6
- Cột "Audio" có thể để trống. Nếu có file mp3, đặt tên file (vd: ni3.mp3) và bỏ vào
  thư mục assets/audio/, rồi ghi tên file đó vào cột Audio.
- Có thể để trống ô nào không có dữ liệu (vd: chưa có Hán Việt) — script sẽ tự bỏ trống trong JSON.
"""

import sys
import json
import os

try:
    import openpyxl
except ImportError:
    print("Thiếu thư viện openpyxl. Cài bằng lệnh: pip install openpyxl")
    sys.exit(1)

# Map: tên cột trong Excel -> tên trường trong JSON
COLUMN_MAP = {
    "Từ/Cụm từ": "hanzi",
    "Pinyin": "pinyin",
    "Hán Việt": "hanviet",
    "Nghĩa tiếng Việt": "nghia",
    "Loại từ": "loai_tu",
    "Ví dụ tiếng Trung": "vidu",
    "Cấp độ HSK": "hsk",
    "Audio": "audio",
}

REQUIRED_COLUMNS = ["Từ/Cụm từ", "Pinyin", "Nghĩa tiếng Việt"]


def convert(input_path: str, output_path: str):
    if not os.path.exists(input_path):
        print(f"Không tìm thấy file: {input_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(input_path, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    header_index = {}
    for idx, name in enumerate(header_row):
        if name is None:
            continue
        clean_name = str(name).strip()
        header_index[clean_name] = idx

    missing = [c for c in REQUIRED_COLUMNS if c not in header_index]
    if missing:
        print("Thiếu cột bắt buộc trong file Excel:", ", ".join(missing))
        print("Các cột tìm thấy:", ", ".join(header_index.keys()))
        sys.exit(1)

    result = []
    row_id = 1
    skipped = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Bỏ qua dòng trống (không có Từ/Cụm từ)
        hanzi_idx = header_index.get("Từ/Cụm từ")
        if hanzi_idx is None or row[hanzi_idx] is None or str(row[hanzi_idx]).strip() == "":
            skipped += 1
            continue

        entry = {"id": row_id}
        for excel_col, json_field in COLUMN_MAP.items():
            idx = header_index.get(excel_col)
            value = row[idx] if idx is not None and idx < len(row) else None

            if json_field == "hsk":
                # Cấp độ HSK phải là số nguyên, mặc định để trống nếu không hợp lệ
                try:
                    value = int(value) if value not in (None, "") else None
                except (ValueError, TypeError):
                    value = None
            else:
                value = str(value).strip() if value not in (None, "") else ""

            entry[json_field] = value

        result.append(entry)
        row_id += 1

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"Đã chuyển đổi {len(result)} từ vựng -> {output_path}")
    if skipped:
        print(f"(Đã bỏ qua {skipped} dòng trống)")


if __name__ == "__main__":
    input_file = sys.argv[1] if len(sys.argv) > 1 else "vocab_data.xlsx"
    output_file = sys.argv[2] if len(sys.argv) > 2 else os.path.join("data", "vocab.json")
    convert(input_file, output_file)
